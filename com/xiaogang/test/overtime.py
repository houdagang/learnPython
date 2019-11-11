import openpyxl
import datetime
import math

work_days=['19-09-02','19-09-03','19-09-04','19-09-05','19-09-06','19-09-09','19-09-10','19-09-11',
'19-09-12','19-09-16','19-09-17','19-09-18','19-09-19','19-09-20','19-09-23','19-09-24','19-09-25','19-09-26','19-09-27','19-09-29',
'19-09-30','19-10-08','19-10-09','19-10-10','19-10-11','19-10-12','19-10-14','19-10-15','19-10-16','19-10-17','19-10-18','19-10-21',
'19-10-22','19-10-23','19-10-24','19-10-25','19-10-28','19-10-29','19-10-30','19-10-31','19-11-01','19-11-04','19-11-05','19-11-06',
'19-11-07','19-11-08','19-11-11','19-11-12','19-11-13','19-11-14','19-11-15','19-11-18','19-11-19','19-11-20','19-11-21',
'19-11-22','19-11-25','19-11-26','19-11-27','19-11-28','19-11-29','19-12-02','19-12-03','19-12-04','19-12-05','19-12-06',
'19-12-09','19-12-10','19-12-11','19-12-12','19-12-13','19-12-16','19-12-17','19-12-18','19-12-19','19-12-20','19-12-23',
'19-12-24','19-12-25','19-12-26','19-12-27','19-12-30','19-12-31']

#总方法
def main():
    user_name = input('请输入要统计的员工名字：')
    flag = selectName(user_name)
    if(flag == False):
        main()
    else:
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = '加班时长'
        sheet['A1'] = '姓名'
        sheet['B1'] = '日期'
        sheet['C1'] = '签到时间（起）'
        sheet['D1'] = '签到时间（止）'
        sheet['E1'] = '时长'
        sheet['F1'] = '合计（h）'
        l = len(flag)
        l1 = l + 1
        all_hours = 0
        for data in flag:
            name = data[0]
            date = data[1]
            ks_date = data[2]
            js_date = data[3]
            jbsc = data[4]
            all_hours += float(jbsc)
            sheet.append([name,date,ks_date,js_date,jbsc])
        sheet['F2'] = all_hours
        sheet.merge_cells('A2:A'+str(l1))
        sheet.merge_cells('F2:F'+str(l1))
        wb.save('D://'+user_name+'加班时长统计.xlsx')
        wb.close()
    flag = input('是否继续？1否,其他任意继续')
    if flag != '1':
        main()

#筛选数据
def selectName(user_name=''):
    list = []
    for i in range(5,max_row):
        obj = []
        name = sheet.cell(i, 1).value
        if name != user_name:
            continue
    
        all_date = sheet.cell(i, 6).value
        date = ''
        #week = ''
        if all_date.strip() != '':
            date = all_date[:9]
            #week = all_date[-3:]
        sb_time = sheet.cell(i,8).value
        sb_zt = sheet.cell(i,9).value
        if not sb_time:
            if sb_zt:
                if sb_zt.strip() == '缺卡':
                    sb_time = '08:30'
        xb_time = sheet.cell(i,10).value
        if bool(sb_time) and bool(xb_time):
            obj.append(name)
            obj.append(date)
            if date.strip() not in work_days:
                obj.append(sb_time)
            else:
                obj.append('17:30')
            obj.append(xb_time)
        else:
            continue
        list.append(obj)
    if len(list) == 0:
        print("查无此人")
        return False
    #print(list)
    user_list = calTime(list)
    return user_list

#计算加班时长
def calTime(user_list):
    return_list = []
    for user in user_list:
        if user[1].strip() in work_days:
            #工作日加班
            day1 = user[1].strip().split('-')
            day2 = user[1].strip().split('-')
            time = user[3].strip().split(':')
            day0 = user[2].strip().split(':')
            day1.extend(time)
            #加班过了凌晨
            today = datetime.datetime.strptime('00' + user[1].strip(),'%Y-%m-%d')
            tomorrow = today + datetime.timedelta(days=1)
            tomo = str(tomorrow).split(' ')[0].split('-')
            time00 = ['8','30']
            day2.append('20')
            day2.append('00')

            tomorrow = int(day2[2]) + 1
            time1 = datetime.datetime(int(day1[0]),int(day1[1]),int(day1[2]),int(day1[3]),int(day1[4]))
            time11 = datetime.datetime(int(tomo[0]),int(tomo[1]),int(tomo[2]),int(day1[3]),int(day1[4]))
            time2 = datetime.datetime(int(day2[0]),int(day2[1]),int(day2[2]),int(day2[3]),int(day2[4]))
            time3 = datetime.datetime(int(tomo[0]),int(tomo[1]),int(tomo[2]),int(time00[0]),int(time00[1]))
            time0 = datetime.datetime(int(day2[0]),int(day2[1]),int(day2[2]),int(day0[0]),int(day0[1]))
            time000 = datetime.datetime(int(tomo[0]),int(tomo[1]),int(tomo[2]),int(0),int(0))
            if time1.__ge__(time2):
                #计算时间差 - 0.5小时
                seconds = (time1 - time0).seconds
                hours = seconds/3600 - 0.5
                jb_hours = math.floor(hours/0.5)*0.5
                #拼接加班时长
                user.append(jb_hours)
                return_list.append(user)
                
                #加班过了凌晨
            elif time11.__le__(time3):
                seconds = (time11 - time000).seconds
                hours = seconds/3600 + 6
                jb_hours = math.floor(hours/0.5)*0.5
                #拼接加班时长
                user.append(jb_hours)
                return_list.append(user)
            else:
                #为满足不计入统计
                continue
        else:
            #节假日加班
            day1 = user[1].strip().split('-')
            day2 = user[1].strip().split('-')
            time1 = user[2].strip().split(':')
            time2 = user[3].strip().split(':')
            day1.extend(time1)
            day2.extend(time2)
            a1 = datetime.datetime(int(day1[0]),int(day1[1]),int(day1[2]),int(day1[3]),int(day1[4]))
            a2 = datetime.datetime(int(day2[0]),int(day2[1]),int(day2[2]),int(day2[3]),int(day2[4]))

            time_sw = '11:30'
            time_xw = '13:00'
            
            sb = datetime.datetime.strptime(user[2].strip(),'%H:%M')
            xb = datetime.datetime.strptime(user[3].strip(),'%H:%M')
            sw = datetime.datetime.strptime(time_sw,'%H:%M')
            xw = datetime.datetime.strptime(time_xw,'%H:%M')

            #计算时间差
            hours = 0
            if sb < sw:
                seconds = (a2 - a1).seconds
                if xb < xw:
                    hours = seconds/3600
                else:
                    hours = seconds/3600 - 1.5
            elif sw <= sb < xw:
                a = datetime.datetime(int(day1[0]),int(day1[1]),int(day1[2]),int('13'),int('00'))
                seconds = (a2 - a).seconds
                hours = seconds/3600
            else:
                seconds = (a2 - a1).seconds
                hours = seconds/3600
            jb_hours = math.floor(hours/0.5)*0.5
            #拼接加班时长
            user.append(jb_hours)
            return_list.append(user)
    return return_list

try:
    global wb
    wb = openpyxl.load_workbook('E://技术中心考勤打卡日统计.xlsx')
except:
    print("未找到文件")
else:
    #获取所有的sheet页的名字
    sheetnames = wb.sheetnames
    #获取第一个sheet页
    sheet = wb[sheetnames[0]]
    #最大行数
    max_row = sheet.max_row + 1
    #最大列数
    #max_col = sheet.max_column + 1
    list = []
    main()
    